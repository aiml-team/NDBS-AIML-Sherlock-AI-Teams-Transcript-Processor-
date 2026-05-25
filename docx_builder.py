"""
docx_builder.py
===============
Renders the filled Customer Profile DOCX by:
  1. Loading word_template.docx via docxtpl (Jinja2 template engine)
  2. Injecting the extracted JSON data directly into the template placeholders
  3. Updating the cover-page client name and date fields

The template already has the exact styling — this module just fills content.
"""

import io
import os
import re
from copy import deepcopy
from datetime import datetime
from docxtpl import DocxTemplate
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ── Internet Research formatting ─────────────────────────────────────────────
# Mirrors the post-process used by the Sherlock_AI_ForAPI renderer: the
# Tavily enrichment block embeds "[Internet Research]" + "• Source:" lines
# into a field's content, which docxtpl renders as one paragraph with soft
# line breaks. We restyle those lines after render and before save.
_IR_HEADER_TEXT = "[Internet Research]"
_IR_RED_HEX = "C00000"
_IR_GREY_HEX = "808080"
_IR_SOURCE_SZ = "16"  # half-points → 8 pt


def _ir_classify(text):
    s = (text or "").strip()
    if s == _IR_HEADER_TEXT:
        return "header"
    if s.startswith("• Source:") or s.startswith("Source:"):
        return "source"
    return "normal"


def _ir_iter_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        yield from _ir_iter_table_paragraphs(table)


def _ir_iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                yield p
            for nested in cell.tables:
                yield from _ir_iter_table_paragraphs(nested)


def _ir_collect_lines(p_el):
    lines = []
    cur_text = ""
    cur_rpr = None
    for r in p_el.findall(qn("w:r")):
        r_rpr = r.find(qn("w:rPr"))
        for child in list(r):
            tag = child.tag
            if tag == qn("w:t"):
                if cur_rpr is None and r_rpr is not None:
                    cur_rpr = r_rpr
                cur_text += (child.text or "")
            elif tag == qn("w:br"):
                lines.append((cur_text, cur_rpr))
                cur_text = ""
                cur_rpr = None
            elif tag == qn("w:tab"):
                cur_text += "\t"
    lines.append((cur_text, cur_rpr))
    return lines


def _ir_apply_header(rpr):
    for tag in ("w:i", "w:iCs", "w:color"):
        for el in rpr.findall(qn(tag)):
            rpr.remove(el)
    rpr.append(OxmlElement("w:i"))
    rpr.append(OxmlElement("w:iCs"))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), _IR_RED_HEX)
    rpr.append(color)


def _ir_apply_source(rpr):
    for tag in ("w:color", "w:sz", "w:szCs"):
        for el in rpr.findall(qn(tag)):
            rpr.remove(el)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), _IR_GREY_HEX)
    rpr.append(color)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), _IR_SOURCE_SZ)
    rpr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), _IR_SOURCE_SZ)
    rpr.append(szCs)


def _ir_restyle_paragraph(p):
    p_el = p._element
    lines = _ir_collect_lines(p_el)
    if not any(_ir_classify(t) != "normal" for t, _ in lines):
        return

    for r in list(p_el.findall(qn("w:r"))):
        p_el.remove(r)

    pPr = p_el.find(qn("w:pPr"))
    insert_idx = list(p_el).index(pPr) + 1 if pPr is not None else 0

    new_elems = []
    for i, (text, base_rpr) in enumerate(lines):
        if i > 0:
            br_run = OxmlElement("w:r")
            if base_rpr is not None:
                br_run.append(deepcopy(base_rpr))
            br_run.append(OxmlElement("w:br"))
            new_elems.append(br_run)
        if not text:
            continue
        cls = _ir_classify(text)
        new_r = OxmlElement("w:r")
        rpr = deepcopy(base_rpr) if base_rpr is not None else OxmlElement("w:rPr")
        if cls == "header":
            _ir_apply_header(rpr)
        elif cls == "source":
            _ir_apply_source(rpr)
        if list(rpr):
            new_r.append(rpr)
        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        new_r.append(t)
        new_elems.append(new_r)

    for el in new_elems:
        p_el.insert(insert_idx, el)
        insert_idx += 1


def _restyle_internet_research(doc):
    for p in _ir_iter_paragraphs(doc):
        _ir_restyle_paragraph(p)


def _safe_content(value) -> str:
    """Return a clean string for any value that might come from the JSON."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, list):
        return "\n".join(str(i) for i in value)
    if isinstance(value, dict):
        if "content" in value:
            return _safe_content(value["content"])
        return "\n".join(f"{k}: {v}" for k, v in value.items())
    return str(value)


def _section(raw_dict: dict) -> dict:
    """Ensure every field in a workstream section has a clean 'content' string."""
    fields = [
        "Current_Processes_Key_Findings",
        "Pain_Points",
        "Proposed_SAP_Solutions_Mapping",
        "Major_Gaps_and_Integrations",
    ]
    out = {}
    for f in fields:
        val = raw_dict.get(f, {})
        content = _safe_content(val.get("content", "") if isinstance(val, dict) else val)
        out[f] = {"content": content or "Not identified in transcripts."}
    return out


def _gbo(raw_dict: dict) -> dict:
    """Build the General_Business_Overview context block."""
    fields = [
        "Schedule_of_Events", "Contacts_Identified", "Industry_Categorization",
        "Revenue_Band", "Legal_Entities_and_Names", "Business_Locations",
        "Fiscal_Year_Format", "Total_SAP_Users", "System_Landscape",
        "Key_Value_Drivers", "Motivations_for_Transformation",
        "Areas_of_Perceived_Competitive_Advantage", "Perceived_Change_Resistance",
        "Technical_Challenges_and_Requirements", "Regulatory_Compliance_Requirements",
        "Transformation_Program_C_Suite_KPIs", "Key_Public_Cloud_Disqualifiers",
    ]
    out = {}
    for f in fields:
        val = raw_dict.get(f, {})
        content = _safe_content(val.get("content", "") if isinstance(val, dict) else val)
        out[f] = {"content": content or "Not identified in transcripts."}
    return out


def build_docx(data: dict, template_path: str | None = None,
               prospect_name: str | None = None) -> bytes:
    """
    Render the template with extracted data and return the DOCX as bytes.

    `data` must match the TEMPLATE_SCHEMA structure from tools.py.
    `template_path` should point to word_template.docx.
    `prospect_name` fills the cover-page {{ prospect_name }} placeholder. If
    not provided, falls back to data['prospect_name'] / data['client_name'].
    """
    if not template_path or not os.path.exists(template_path):
        raise FileNotFoundError(
            f"word_template.docx not found at: {template_path}\n"
            "Place word_template.docx in the same folder as main.py."
        )

    tpl = DocxTemplate(template_path)

    # Resolve prospect_name from the explicit kwarg, then from the data dict
    # (supports both new "prospect_name" and legacy "client_name" keys).
    resolved_prospect = (
        _safe_content(prospect_name)
        or _safe_content(data.get("prospect_name", ""))
        or _safe_content(data.get("client_name", ""))
        or "Client Name"
    )
    doc_date = _safe_content(data.get("document_date", "")) or \
               datetime.now().strftime("%d %B %Y")

    # Build the full Jinja2 context — mirrors {{ data.SECTION.FIELD.content }}
    context = {
        "prospect_name": resolved_prospect,
        "data": {
            "General_Business_Overview": _gbo(
                data.get("General_Business_Overview", {})),

            "Idea_to_Market": _section(
                data.get("Idea_to_Market", {})),
            "Source_to_Pay_S2P": _section(
                data.get("Source_to_Pay_S2P", {})),
            "Plan_to_Produce_P2P": _section(
                data.get("Plan_to_Produce_P2P", {})),
            "Detect_to_Correct_D2C": _section(
                data.get("Detect_to_Correct_D2C", {})),
            "Forecast_to_Fulfill_F2F": _section(
                data.get("Forecast_to_Fulfill_F2F", {})),
            "Warehouse_Execution_WM_EWM": _section(
                data.get("Warehouse_Execution_WM_EWM", {})),
            "Lead_to_Cash_L2C": _section(
                data.get("Lead_to_Cash_L2C", {})),
            "Logistics_Planning_and_Transportation_TM": _section(
                data.get("Logistics_Planning_and_Transportation_TM", {})),
            "Request_to_Service_R2S": _section(
                data.get("Request_to_Service_R2S", {})),
            "Record_to_Report_R2R": _section(
                data.get("Record_to_Report_R2R", {})),
            "Acquire_to_Dispose_A2D": _section(
                data.get("Acquire_to_Dispose_A2D", {})),
            "Environmental_Social_and_Governance_ESG_Processes": _section(
                data.get("Environmental_Social_and_Governance_ESG_Processes", {})),
            "Hire_to_Retire_H2R": _section(
                data.get("Hire_to_Retire_H2R", {})),
            "Enterprise_Reporting_Data_and_Analytics_Strategy": _section(
                data.get("Enterprise_Reporting_Data_and_Analytics_Strategy", {})),
        }
    }

    tpl.render(context)

    # ── Patch remaining cover-page fields (date and "FirstName LastName"). ──
    # The "Client's full name" placeholder is now a real Jinja template var
    # ({{ prospect_name }}) and is filled by tpl.render(context) above, so we
    # only need to fix up the date field's cached value and the author line.
    _patch_cover(tpl, doc_date)

    # Restyle [Internet Research] header (italic + red) and • Source: bullets
    # (8 pt + grey) inserted by the Tavily enrichment.
    try:
        _restyle_internet_research(tpl.docx)
    except Exception as restyle_err:
        print(f"⚠ Internet Research restyle skipped: {restyle_err}")

    buf = io.BytesIO()
    tpl.save(buf)
    return buf.getvalue()


def _patch_cover(tpl: DocxTemplate, doc_date: str):
    """
    Patch the cover page date field and "FirstName LastName" placeholder.
    Operates directly on the rendered XML.
    """
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    body = tpl.docx.element.body

    # Walk all paragraphs in tables (cover is a table)
    for tbl in body.iter("{%s}tbl" % W):
        for p in tbl.iter("{%s}p" % W):
            # Collect full text of paragraph
            texts = [r.text or "" for r in p.iter("{%s}t" % W)]
            full  = "".join(texts)

            # Replace "FirstName LastName"
            if full.strip() == "FirstName LastName":
                _replace_para_text(p, W, "")   # leave blank or put author

            # Replace the date field text (keeps the field, updates cached value)
            elif re.search(r"\d{1,2}\s+\w+\s+\d{4}", full):
                # Update the cached date text inside the field
                for t in p.iter("{%s}t" % W):
                    if t.text and re.search(r"\d{1,2}\s+\w+\s+\d{4}", t.text):
                        t.text = doc_date
                        break


def _replace_para_text(para, W: str, new_text: str):
    """Replace all <w:t> text runs in a paragraph with a single new text run."""
    runs = list(para.iter("{%s}r" % W))
    if not runs:
        return
    # Clear all runs except first, set first run's text
    first_t = runs[0].find("{%s}t" % W)
    if first_t is None:
        from lxml import etree
        first_t = etree.SubElement(runs[0], "{%s}t" % W)
    first_t.text = new_text
    if new_text and (new_text.startswith(" ") or new_text.endswith(" ")):
        first_t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    # Remove subsequent runs
    for r in runs[1:]:
        r.getparent().remove(r)